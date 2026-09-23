"""DOCX generation and parsing for the localization review table.

No ARB semantics here beyond the Row/section shape from arb.py: this module
takes rows in, emits a table; reads a table, emits rows out.

Row identity: a visible, de-emphasized 5th "Key" column holds
"<arb key>|<sha1(english)[:8]>". This was originally a hidden-text
(w:vanish) run instead of a real column, which is cleaner on the page --
but Google Docs' .docx importer silently drops hidden-text runs on upload,
losing row identity entirely. A real column survives every editor because
it's just ordinary cell content.
"""
import hashlib
import itertools
from typing import List, NamedTuple, Optional, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import arb
import validate

MARKER_SEP = "|"
HEADER = ["English", "Current", "Proposed", "Comments", "Key (do not edit)"]
COL_WIDTHS_CM = [7.8, 5.8, 5.8, 3.4, 2.6]
KEY_COLUMN = 4
CJK_FONT = "Malgun Gothic"
LANG_TAGS = {"ko": "ko-KR", "de": "de-DE", "en": "en-US", "tr": "tr-TR"}
KEY_GRAY = RGBColor(0x99, 0x99, 0x99)


class ImportedRow(NamedTuple):
    key: Optional[str]
    marker_checksum: Optional[str]
    visible_english: str
    current: str
    proposed: Optional[str]
    comment: Optional[str]


def checksum(text: str) -> str:
    """Hashed over stripped text: several editors (confirmed: LibreOffice's
    save cycle) trim leading/trailing whitespace from a run even with zero
    user edits, and some real ARB strings intentionally end in a space
    (e.g. "Your " for concatenation). Stripping keeps the identity check
    sensitive to real content changes without false-rejecting on that."""
    return hashlib.sha1(text.strip().encode("utf-8")).hexdigest()[:8]


def _set_landscape(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.makeelement(qn("w:tblHeader"), {})
    tr_pr.append(header)


def _set_run_fonts(run, cjk_name: str = CJK_FONT) -> None:
    """Explicit per-run font, defense in depth against theme resolution."""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), cjk_name)
    rfonts.set(qn("w:cs"), cjk_name)


def _fix_default_fonts(doc: Document, locale: str, cjk_name: str = CJK_FONT) -> None:
    """python-docx's default template ships blank East Asian theme fonts
    (theme1.xml <a:ea typeface=""/>), and every run inherits that via
    docDefaults -> eastAsiaTheme=minorEastAsia. A strict OOXML renderer
    resolving that chain gets an empty font name and can render nothing.
    Replace the theme references with explicit font names so this holds
    for exported text AND for whatever the translator types afterward."""
    styles_elm = doc.styles.element
    doc_defaults = styles_elm.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return

    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        return
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme", "w:cstheme"):
            if qn(attr) in rfonts.attrib:
                del rfonts.attrib[qn(attr)]
        rfonts.set(qn("w:ascii"), "Calibri")
        rfonts.set(qn("w:hAnsi"), "Calibri")
        rfonts.set(qn("w:eastAsia"), cjk_name)
        rfonts.set(qn("w:cs"), cjk_name)
    lang = rpr.find(qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:eastAsia"), LANG_TAGS.get(locale, locale))


def _write_key_cell(cell, key: str, english: str) -> None:
    cell.text = MARKER_SEP.join([key, checksum(english)])
    run = cell.paragraphs[0].runs[0]
    run.font.size = Pt(6)
    run.font.color.rgb = KEY_GRAY


def write(path, module: str, locale: str, rows: List["arb.Row"]) -> None:
    doc = Document()
    doc.core_properties.subject = "%s|%s" % (module, locale)
    doc.core_properties.title = "StudyU %s localization review \u2014 %s" % (module, locale)
    _fix_default_fonts(doc, locale)
    _set_landscape(doc)

    doc.add_heading("StudyU %s \u2014 %s localization review" % (module, locale), level=1)
    doc.add_paragraph(
        "Fill in the Proposed column only where you would change the current "
        "translation. Leave it blank to accept the current translation as-is. "
        "Use Comments for anything you want us to know \u2014 tone, context, ambiguity. "
        "Where English lists placeholders (e.g. {name}), your proposed text must "
        "contain the exact same placeholders. Please don't edit the small grey "
        "Key column on the right \u2014 it's how your edits get matched back to "
        "the right string."
    )

    for section, section_rows in itertools.groupby(rows, key=lambda r: r.section):
        if section:
            doc.add_heading(section, level=2)
        table = doc.add_table(rows=1, cols=len(HEADER))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _repeat_header(table.rows[0])
        for cell, label, width in zip(table.rows[0].cells, HEADER, COL_WIDTHS_CM):
            cell.text = label
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            _set_run_fonts(run)
            cell.width = Cm(width)

        for row in section_rows:
            cells = table.add_row().cells
            for cell, width in zip(cells, COL_WIDTHS_CM):
                cell.width = Cm(width)

            english_p = cells[0].paragraphs[0]
            english_run = english_p.add_run(row.english)
            _set_run_fonts(english_run)
            placeholders = validate.extract_placeholders(row.english)
            if placeholders:
                hint = cells[0].add_paragraph()
                hint_run = hint.add_run("Placeholders: " + ", ".join(sorted(set(placeholders))))
                hint_run.italic = True
                hint_run.font.size = Pt(8)
                _set_run_fonts(hint_run)

            cells[1].text = row.current or ""
            if row.current:
                _set_run_fonts(cells[1].paragraphs[0].runs[0])
            cells[2].text = ""
            cells[3].text = ""
            _write_key_cell(cells[KEY_COLUMN], row.key, row.english)

    doc.save(str(path))


def _visible_text(cell) -> str:
    if not cell.paragraphs:
        return ""
    return "".join(run.text for run in cell.paragraphs[0].runs)


HEADING_RE = "^StudyU (\\S+) \u2014 (\\S+) localization review$"


def _module_locale_from_heading(doc: Document) -> Optional[Tuple[str, str]]:
    import re

    for p in doc.paragraphs:
        m = re.match(HEADING_RE, p.text.strip())
        if m:
            return m.group(1), m.group(2)
    return None


def read(path) -> Tuple[str, str, List[ImportedRow]]:
    """Google Docs' .docx export wipes custom core properties (confirmed:
    Subject and Title both come back blank), so those can't be trusted as
    the sole source of module/locale. The Heading 1 title survives as
    ordinary paragraph text in every editor tested, so it's the primary
    source; core properties remain a fallback for documents that never
    left Word."""
    doc = Document(str(path))
    detected = _module_locale_from_heading(doc)
    if detected is None:
        subject = doc.core_properties.subject or ""
        if MARKER_SEP not in subject:
            raise ValueError(
                "could not determine module/locale: no 'StudyU <module> \u2014 <locale> "
                "localization review' heading found, and the Subject property is empty "
                "\u2014 was this document exported by this tool?"
            )
        detected = tuple(subject.split(MARKER_SEP, 1))
    module, locale = detected

    rows: List[ImportedRow] = []
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells[: len(HEADER)]]
        if header != HEADER:
            continue
        for tr in table.rows[1:]:
            cells = tr.cells
            marker = cells[KEY_COLUMN].text.strip()
            key = mchecksum = None
            if marker and MARKER_SEP in marker:
                key, mchecksum = marker.rsplit(MARKER_SEP, 1)
            rows.append(
                ImportedRow(
                    key=key,
                    marker_checksum=mchecksum,
                    visible_english=_visible_text(cells[0]),
                    current=cells[1].text.strip(),
                    proposed=(cells[2].text.strip() or None),
                    comment=(cells[3].text.strip() or None),
                )
            )
    return module, locale, rows
