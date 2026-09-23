import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import arb  # noqa: E402
import document  # noqa: E402
import l10n_docx  # noqa: E402
from docx import Document  # noqa: E402


def _copy_fixtures(tmp_path):
    base = tmp_path / "arb"
    base.mkdir()
    for f in arb.FIXTURES.glob("*.arb"):
        shutil.copy(f, base / f.name)
    return base


def _row_index(rows, key):
    return 1 + [r.key for r in rows].index(key)  # +1 for header row


def _build_returned_docx(tmp_path, base):
    rows = arb.rows("app", "ko", base=base)
    exported = tmp_path / "exported.docx"
    document.write(exported, "app", "ko", rows)

    doc = Document(str(exported))
    table = doc.tables[0]

    # 1. valid proposal, placeholder preserved
    r = table.rows[_row_index(rows, "greeting")]
    r.cells[2].text = "\uc548\ub155, {name}!"

    # 2. no-op (blank proposed) with a comment
    r = table.rows[_row_index(rows, "loading")]
    r.cells[3].text = "too formal for teens"

    # 3. placeholder violation: drops {total}
    r = table.rows[_row_index(rows, "studies_count_total")]
    r.cells[2].text = "\uac74\uc218\uac00 \uc788\uc2b5\ub2c8\ub2e4"

    # 4. edited English cell -> checksum mismatch
    r = table.rows[_row_index(rows, "brace_literal")]
    visible_run = r.cells[0].paragraphs[0].runs[0]
    visible_run.text = "Something else entirely"
    r.cells[2].text = "\ubb34\uc5b8\uac00"

    # 5. duplicated row: clone the (unedited) "loading" row's key onto a new row
    loading_row = table.rows[_row_index(rows, "loading")]
    dup = table.add_row()
    dup.cells[0].text = "Loading"
    document._write_key_cell(dup.cells[document.KEY_COLUMN], "loading", "Loading")
    dup.cells[1].text = loading_row.cells[1].text
    dup.cells[2].text = "\ub2e4\ub978 \uc81c\uc548"

    returned = tmp_path / "returned.docx"
    doc.save(str(returned))
    return returned


def test_end_to_end_applies_valid_and_reports_everything_else(tmp_path):
    base = _copy_fixtures(tmp_path)
    returned = _build_returned_docx(tmp_path, base)

    report = l10n_docx.import_docx(returned, base=base, expected_module="app", expected_locale="ko")

    assert dict(report.applied) == {"greeting": "\uc548\ub155, {name}!"}

    reasons = {s.key: s.reason for s in report.skipped}
    assert "placeholder validation failed" in reasons["studies_count_total"]
    assert "row edited" in reasons["brace_literal"]
    # both copies of the duplicated key are rejected, not just the clone
    dup_reasons = [s for s in report.skipped if s.key == "loading"]
    assert len(dup_reasons) == 2
    assert all("duplicated row" in s.reason for s in dup_reasons)

    assert report.comments == []  # the commented row (loading) was a duplicate, so it's skipped, not applied+commented


def test_written_arb_has_only_the_valid_proposal_applied(tmp_path):
    base = _copy_fixtures(tmp_path)
    returned = _build_returned_docx(tmp_path, base)
    l10n_docx.import_docx(returned, base=base, expected_module="app", expected_locale="ko")

    written = arb.load("app", "ko", base=base)
    assert written["greeting"] == "\uc548\ub155, {name}!"
    assert written["studies_count_total"] == "{total, plural, =1{1 study} other{{total} studies}}"
    assert written["brace_literal"] == "\uc911\uad04\ud638 {braces} \uc720\uc9c0\ub428"
    assert written["loading"] == "\ub85c\ub529 \uc911"


def test_reimporting_with_zero_valid_proposals_is_a_pure_noop(tmp_path):
    base = _copy_fixtures(tmp_path)
    rows = arb.rows("app", "ko", base=base)
    exported = base / "exported.docx"
    document.write(exported, "app", "ko", rows)

    before = (base / "app_ko.arb").read_bytes()
    report = l10n_docx.import_docx(exported, base=base, expected_module="app", expected_locale="ko")
    after = (base / "app_ko.arb").read_bytes()

    assert report.applied == []
    assert report.skipped == []
    assert before == after


def test_wrong_module_argument_aborts_whole_import(tmp_path):
    import pytest

    base = _copy_fixtures(tmp_path)
    rows = arb.rows("app", "ko", base=base)
    exported = base / "exported.docx"
    document.write(exported, "app", "ko", rows)

    with pytest.raises(ValueError):
        l10n_docx.import_docx(exported, base=base, expected_module="designer", expected_locale="ko")
