import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import arb  # noqa: E402
import document  # noqa: E402


def _generated(tmp_path):
    rows = arb.rows("app", "ko", base=arb.FIXTURES)
    path = tmp_path / "fixture.docx"
    document.write(path, "app", "ko", rows)
    return path, rows


def test_module_locale_marker_round_trips(tmp_path):
    path, _ = _generated(tmp_path)
    module, locale, _ = document.read(path)
    assert (module, locale) == ("app", "ko")


def test_every_row_key_round_trips(tmp_path):
    path, rows = _generated(tmp_path)
    _, _, imported = document.read(path)
    assert [r.key for r in imported] == [r.key for r in rows]


def test_checksum_matches_visible_english_on_export(tmp_path):
    path, rows = _generated(tmp_path)
    _, _, imported = document.read(path)
    for row, imp in zip(rows, imported):
        assert imp.visible_english == row.english
        assert imp.marker_checksum == document.checksum(row.english)


def test_current_translation_visible_and_proposed_blank(tmp_path):
    path, rows = _generated(tmp_path)
    _, _, imported = document.read(path)
    by_key = {r.key: r for r in imported}
    assert by_key["loading"].current == "\ub85c\ub529 \uc911"
    assert by_key["loading"].proposed is None
    # missing-in-target key exports with an empty Current cell
    assert by_key["only_in_en"].current == ""

def test_checksum_ignores_incidental_whitespace_trimming():
    """Some real ARB strings intentionally end in a space (e.g. "Your " for
    concatenation), and at least one real editor's save cycle (LibreOffice,
    confirmed) trims trailing whitespace from a run with zero user edits.
    The identity check must not false-reject those rows."""
    assert document.checksum("Your ") == document.checksum("Your")
    assert document.checksum(" and ") == document.checksum("and")
    assert document.checksum("Hi") != document.checksum("Hi!")

def test_key_column_is_a_real_visible_cell_not_hidden_text(tmp_path):
    """The identity marker was originally hidden text (w:vanish), but
    Google Docs' .docx importer silently drops hidden runs on upload,
    losing row identity. It must now be ordinary, findable cell text."""
    path, rows = _generated(tmp_path)
    from docx import Document as _Doc

    doc = _Doc(str(path))
    table = doc.tables[0]
    key_cell = table.rows[1].cells[document.KEY_COLUMN]
    run = key_cell.paragraphs[0].runs[0]
    assert run.font.hidden is not True
    assert run.text == "loading|" + document.checksum("Loading")
    # de-emphasized, not literally invisible
    assert run.font.size is not None and run.font.size.pt <= 8


def test_placeholder_hint_present_for_keys_with_placeholders(tmp_path):
    path, rows = _generated(tmp_path)
    from docx import Document as _Doc

    doc = _Doc(str(path))
    table = doc.tables[0]
    by_key = {r.key: i for i, r in enumerate(rows)}
    greeting_cell = table.rows[1 + by_key["greeting"]].cells[0]
    text = "\n".join(p.text for p in greeting_cell.paragraphs)
    assert "Placeholders: name" in text


def test_section_heading_present(tmp_path):
    path, _ = _generated(tmp_path)
    from docx import Document as _Doc

    doc = _Doc(str(path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "General" in headings
